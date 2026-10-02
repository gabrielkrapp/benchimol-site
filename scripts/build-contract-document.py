"""Build the editable Benchimol contract from its local Markdown source.

Uses python-docx from the bundled Codex workspace runtime. Rendering and visual
verification are performed separately with the documents skill renderer.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/contracts/contrato-benchimol.md"
OUTPUT = ROOT / "docs/contracts/contrato-benchimol.docx"


def rich_text(paragraph, value):
    """The source uses only bold emphasis; leave URLs and placeholders literal."""
    parts = re.split(r"(\*\*.*?\*\*)", value)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run(part[2:-2] if part.startswith("**") else part)
        if part.startswith("**"):
            run.bold = True


def cell_margins(cell, top=100, start=120, bottom=100, end=120):
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    cell._tc.get_or_add_tcPr().append(margins)


def terms_table(document, rows):
    table = document.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(2.05), Inches(4.75))
    for column, width in zip(table.columns, widths):
        column.width = width
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D5D5D5")
        borders.append(element)
    table._tbl.tblPr.append(borders)
    for i, row_values in enumerate(rows):
        row = table.rows[i]
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        if i == 0:
            repeat = OxmlElement("w:tblHeader")
            row._tr.get_or_add_trPr().append(repeat)
        for j, value in enumerate(row_values):
            cell = row.cells[j]
            cell.width = widths[j]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            rich_text(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
                if i == 0:
                    run.bold = True
            if i == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EEEEEE")
                cell._tc.get_or_add_tcPr().append(shade)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_page_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, code, separate, fallback, end):
        run._r.append(element)


def main():
    document = Document()
    # The bundled default template contains accent rules in paragraph styles.
    # This contract uses plain black headings with no decorative borders.
    for border in list(document.styles.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.04
    normal.paragraph_format.widow_control = True
    normal._element.get_or_add_rPr().append(OxmlElement("w:lang"))
    normal._element.rPr.find(qn("w:lang")).set(qn("w:val"), "pt-BR")
    for name, size in (("Title", 15), ("Heading 1", 12), ("Heading 2", 12)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(7 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.text = "Clínica de Olhos Benchimol • Prestação de serviços"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Página ")
    add_page_field(footer, "PAGE")
    footer.add_run(" de ")
    add_page_field(footer, "NUMPAGES")
    for run in footer.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0, 0, 0)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    i = 0
    title_seen = False
    signature_section = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "<!-- PAGEBREAK -->":
            document.add_page_break()
            signature_section = False
        elif line.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                values = [value.strip() for value in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"[:\- ]+", value) for value in values):
                    rows.append(values)
                i += 1
            terms_table(document, rows)
            continue
        elif line.startswith("# "):
            paragraph = document.add_paragraph(style="Title" if not title_seen else "Heading 1")
            rich_text(paragraph, line[2:])
            title_seen = True
        elif line.startswith("## "):
            text = line[3:]
            signature_section = text == "Assinaturas"
            paragraph = document.add_paragraph(style="Heading 1")
            rich_text(paragraph, text)
        else:
            paragraph = document.add_paragraph()
            rich_text(paragraph, line)
            if signature_section:
                paragraph.paragraph_format.space_after = Pt(6)
                paragraph.paragraph_format.keep_with_next = not line.startswith("Assinatura:")
            if line.startswith("**MINUTA"):
                paragraph.paragraph_format.space_after = Pt(10)
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
        i += 1

    properties = document.core_properties
    properties.title = "Contrato de prestação de serviços — Clínica de Olhos Benchimol"
    properties.subject = "Migração WordPress/Elementor para Next.js e painel administrativo"
    properties.author = "Gabriel Krapp Oliveira"
    properties.keywords = "minuta, migração, Next.js, blog, painel administrativo, SEO"
    properties.comments = "Versão simplificada com dados confirmados pelo contratante do documento."
    document.save(OUTPUT)
    print(f"DOCX criado: {OUTPUT}")


if __name__ == "__main__":
    main()
